#!/usr/bin/env python3
"""Role test cases as a workbook testers fill in (#4804).

Generated from the code so it never drifts from what each role can do:
the menu (apps/web/src/config/navigation.ts, same visibility rule as the
sidebar), the role grants (crates/medbrains-core/src/access/roles.rs) and the
permission codes (crates/medbrains-core/src/permissions.rs). The day-to-day
jobs come from scripts/role_test_workflows.py.

    python3 scripts/generate_role_test_cases.py
    -> docs/testing/MedBrains_Role_Test_Cases.xlsx

Per role: sign in; every screen the role should open; one case listing what
it must NOT see plus direct-URL checks; then its workflows. Re-generate
after a permission or menu change; results already typed in are not kept.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
from role_test_workflows import WORKFLOWS  # noqa: E402

OUT = ROOT / "docs" / "testing" / "MedBrains_Role_Test_Cases.xlsx"
ADMIN_ROLES = {"super_admin", "hospital_admin"}
# Probed by URL when a role cannot see them: the screens a leak would hurt most.
SENSITIVE = ["/admin/users", "/audit", "/billing", "/mrd", "/pharmacy", "/lab", "/ipd", "/hr/payroll"]
HEADER = ["ID", "Type", "Area", "What to test", "Steps", "Expected", "Permission",
          "Result", "Tester", "Date", "Notes"]
WIDTHS = [11, 13, 18, 34, 60, 60, 30, 11, 14, 12, 30]
RESULTS = '"Pass,Fail,Blocked,Not run"'


def permission_codes() -> dict[str, str]:
    """`module::sub::CONST` path → "module.sub.code"."""
    stack: list[str] = []
    paths: dict[str, str] = {}
    for line in (ROOT / "crates/medbrains-core/src/permissions.rs").read_text().splitlines():
        if m := re.match(r"\s*pub mod (\w+) \{", line):
            stack.append(m.group(1))
        elif m := re.match(r'\s*pub const (\w+): &str = "([^"]+)";', line):
            paths["::".join([*stack, m.group(1)])] = m.group(2)
        elif re.match(r"\s*\}\s*$", line) and stack:
            stack.pop()
    return paths


def roles(paths: dict[str, str]) -> list[tuple[str, str, set[str]]]:
    text = (ROOT / "crates/medbrains-core/src/access/roles.rs").read_text()
    out = []
    for block in re.split(r"\n\s*BuiltInRole \{", text)[1:]:
        code = re.search(r'code: "([a-z_]+)"', block).group(1)
        name = re.search(r'name: "([^"]+)"', block).group(1)
        granted = {paths[p] for p in re.findall(r"permissions::([a-z_0-9:]+::[A-Z_0-9]+)", block) if p in paths}
        out.append((code, name, granted))
    return out


def menu() -> list[dict]:
    nav = (ROOT / "apps/web/src/config/navigation.ts").read_text()
    nav = nav[nav.index("export const NAV_GROUPS"):]
    labels = json.loads((ROOT / "apps/web/public/locales/en/nav.json").read_text())
    group, items = "", []
    pattern = r'key: "([a-zA-Z_]+)",\s*\n\s*items:|i18nKey: "(\w+)",\s*\n\s*path: "([^"]+)",(.*?)(?=\n\s*\},|\n\s*children:)'
    for m in re.finditer(pattern, nav, re.S):
        if m.group(1):
            group = m.group(1)
            continue
        rest = m.group(4)
        one = re.search(r'requiredPermission: "([^"]+)"', rest)
        many = re.search(r"requiredPermissions: \[([^\]]*)\]", rest)
        items.append({
            "group": labels.get("group" + group[:1].upper() + group[1:], group),
            "label": labels.get(m.group(2), m.group(2)),
            "path": m.group(3),
            "one": one.group(1) if one else None,
            "any": re.findall(r'"([^"]+)"', many.group(1)) if many else [],
            "module": "moduleCode" in rest,
        })
    return items


def visible(item: dict, granted: set[str]) -> bool:
    """The sidebar's rule (AppLayout): the single code, and any one of the list."""
    return (item["one"] is None or item["one"] in granted) and (
        not item["any"] or any(p in granted for p in item["any"]))


def cases_for(code: str, name: str, granted: set[str], items: list[dict]) -> list[list[str]]:
    rows: list[list[str]] = []

    def add(kind: str, area: str, what: str, steps: str, expected: str, perm: str = "") -> None:
        rows.append([f"{code[:6].upper()}-{len(rows) + 1:03d}", kind, area, what, steps, expected, perm,
                     "Not run", "", "", ""])

    login = f"test_{code}"[:30]
    add("Sign in", "Sign in", f"Sign in as {name}",
        f"app.medbrains.in → hospital code AMH → username {login} and the password from your test-accounts file.",
        f"You are signed in; the header shows 'TEST - {name}'; you land on a page you can use (not an error).")
    seen = [i for i in items if visible(i, granted)]
    for item in seen:
        add("Should see", item["group"], f"Open {item['label']}",
            f"Menu → {item['group']} → {item['label']}  ({item['path']})",
            "The screen opens with no 'access denied' or error message. Lists load; an empty list says so. "
            "Buttons to add or change things appear only where this role should do that."
            + (" (Hidden if the hospital has this module switched off.)" if item["module"] else ""),
            item["one"] or ", ".join(item["any"][:3]) + ("…" if len(item["any"]) > 3 else ""))
    hidden = [i for i in items if not visible(i, granted)]
    if hidden:
        add("Must not see", "Menu", "Menu shows only this role's screens",
            "Look through every menu group.",
            "These must NOT be in the menu: " + ", ".join(i["label"] for i in hidden) + ".")
        by_path = {i["path"]: i for i in hidden}
        for path in [p for p in SENSITIVE if p in by_path][:4]:
            add("Must not see", by_path[path]["group"], f"Typing the address of {by_path[path]['label']}",
                f"Type app.medbrains.in{path} into the address bar.",
                "You get an 'access denied' page or are sent back — none of that screen's data is shown, even briefly.",
                by_path[path]["one"] or "")
    for w in WORKFLOWS.get(code, []):
        add("Workflow", w["area"], w["title"], w["steps"], w["expected"])
    return rows


def style_sheet(ws, rows: list[list[str]]) -> None:
    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4332")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1=RESULTS, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"H2:H{len(rows) + 1}")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:K{len(rows) + 1}"


def readme(ws, total: int, n_roles: int) -> None:
    lines = [
        ("MedBrains — role test cases", True),
        ("", False),
        ("Why: signed in as an administrator you see every screen. These cases check what each role really sees and does.", False),
        (f"{n_roles} roles, {total} cases. One tab per role; the Summary tab totals results.", False),
        ("", False),
        ("How to test", True),
        ("1. Take your role's username and password from the test-accounts file (shared via the password manager).", False),
        ("2. Sign in at app.medbrains.in with hospital code AMH.", False),
        ("3. Work down your role's tab. Set Result to Pass, Fail, Blocked (could not test) or leave Not run. Add your name and date.", False),
        ("4. For every Fail: write what you saw in Notes, take a screenshot, and report it with the case ID.", False),
        ("", False),
        ("Case types", True),
        ("Sign in — the account works.", False),
        ("Should see — every screen this role's permissions open.", False),
        ("Must not see — screens hidden from this role, and refused even when the address is typed.", False),
        ("Workflow — the role's everyday jobs, end to end. 'as test_doctor' means another role does that step.", False),
        ("", False),
        ("Careful: this is the live hospital. Use test patients you create (name them 'TEST …'). Do not change real patients' records.", True),
        ("Generated from the code by scripts/generate_role_test_cases.py — regenerate after permission or menu changes.", False),
    ]
    for text, bold in lines:
        ws.append([text])
        ws.cell(ws.max_row, 1).font = Font(bold=bold, size=14 if text.startswith("MedBrains") else 11)
    ws.column_dimensions["A"].width = 120


def main() -> int:
    paths = permission_codes()
    items = menu()
    wb = Workbook()
    readme_ws = wb.active
    readme_ws.title = "Read me"
    summary = wb.create_sheet("Summary")
    summary.append(["Role", "Login", "Permissions", "Cases", "Pass", "Fail", "Blocked", "Not run"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    total = 0
    tested = [r for r in roles(paths) if r[0] not in ADMIN_ROLES]
    for code, name, granted in tested:
        rows = cases_for(code, name, granted, items)
        total += len(rows)
        title = name[:31]
        style_sheet(wb.create_sheet(title), rows)
        ref = f"'{title}'!H2:H{len(rows) + 1}"
        r = summary.max_row + 1
        summary.append([name, f"test_{code}"[:30], len(granted), len(rows),
                        f'=COUNTIF({ref},"Pass")', f'=COUNTIF({ref},"Fail")',
                        f'=COUNTIF({ref},"Blocked")', f'=COUNTIF({ref},"Not run")'])
        summary.cell(r, 1).hyperlink = f"#'{title}'!A1"
    for col, width in zip("ABCDEFGH", [28, 32, 12, 8, 8, 8, 9, 9]):
        summary.column_dimensions[col].width = width
    readme(readme_ws, total, len(tested))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"{len(tested)} roles, {total} cases -> {OUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
